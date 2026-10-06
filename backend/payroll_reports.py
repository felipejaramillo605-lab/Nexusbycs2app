"""Reportes de nomina: Excel mensual de gastos de personal (tipo reporte) y colilla de pago en PDF.

Plantilla por defecto: fondo blanco y colorimetria de Nexus. La colilla muestra el logo que el negocio haya definido.
"""

from __future__ import annotations

from io import BytesIO
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from payroll_co import DISCLAIMER

PRIMARY_HEX = "7C3AED"
PRIMARY_DARK_HEX = "6D28D9"
SOFT_HEX = "F3EEFF"
BORDER_HEX = "E4E7EC"
TEXT_HEX = "101828"
MUTED_HEX = "667085"

MONTHS_ES = [
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
]
STATUS_ES = {"draft": "Borrador", "approved": "Aprobada", "paid": "Pagada", "cancelled": "Cancelada"}
FREQUENCY_ES = {"monthly": "Mensual", "biweekly": "Quincenal"}
COP = '"$" #,##0;[Red]-"$" #,##0'


def period_label(run: dict) -> str:
    month = MONTHS_ES[int(run["month"]) - 1].capitalize()
    if run.get("frequency") == "biweekly":
        return f"{'Primera' if run.get('half') == 1 else 'Segunda'} quincena de {month.lower()} de {run['year']}"
    return f"{month} de {run['year']}"


def fmt(value) -> str:
    return "$ " + f"{float(value or 0):,.0f}".replace(",", ".")


def _sum_codes(lines: List[dict], key: str) -> dict:
    """Suma por concepto (code+label) de un grupo de rubros entre todos los trabajadores."""
    totals: dict = {}
    for line in lines:
        for row in (line.get("computed") or {}).get(key, []):
            label = row["label"].split(" — ")[0]
            totals[(row["code"], label)] = totals.get((row["code"], label), 0.0) + float(row["amount"])
    return totals


# ----------------------------------------------------------------------------------------------- Excel


def build_payroll_workbook(run: dict, organization_name: str) -> bytes:
    thin = Side(style="thin", color=BORDER_HEX)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor=PRIMARY_HEX)
    soft_fill = PatternFill("solid", fgColor=SOFT_HEX)
    head_font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    base_font = Font(name="Calibri", size=10, color=TEXT_HEX)
    bold = Font(name="Calibri", size=10, bold=True, color=TEXT_HEX)
    wb = Workbook()
    salaried = [x for x in run.get("lines", []) if x.get("contract_type") == "fixed_salary" and x.get("computed")]
    commission = [x for x in run.get("lines", []) if x.get("contract_type") == "service_commission"]

    def title(ws, text, subtitle, width):
        ws.sheet_view.showGridLines = False
        ws["A1"] = organization_name
        ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=PRIMARY_DARK_HEX)
        ws["A2"] = text
        ws["A2"].font = Font(name="Calibri", size=12, bold=True, color=TEXT_HEX)
        ws["A3"] = subtitle
        ws["A3"].font = Font(name="Calibri", size=10, color=MUTED_HEX)
        for col in range(1, width + 1):
            ws.cell(row=4, column=col).fill = PatternFill("solid", fgColor=PRIMARY_HEX)
        ws.row_dimensions[4].height = 3

    def header_row(ws, row, labels):
        for index, label in enumerate(labels, start=1):
            cell = ws.cell(row=row, column=index, value=label)
            cell.fill, cell.font, cell.border = head_fill, head_font, border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 30

    def body_cell(ws, row, col, value, money=False, strong=False, fill=None):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font, cell.border = (bold if strong else base_font), border
        if money:
            cell.number_format = COP
            cell.alignment = Alignment(horizontal="right")
        if fill:
            cell.fill = fill
        return cell

    # ---------------- Resumen
    ws = wb.active
    ws.title = "Resumen"
    status_label = STATUS_ES.get(run.get("status"), run.get("status"))
    subtitle = f"Periodo: {period_label(run)}  ·  Estado: {status_label}  ·  Versión {run.get('version', 1)}"
    title(ws, "Reporte de gastos de personal", subtitle, 3)
    totals = run.get("totals", {})
    kpis = [
        ("Devengado total (empleados con contrato fijo)", totals.get("gross", 0)),
        ("Deducciones a cargo de los empleados", totals.get("deductions", 0)),
        ("Neto a pagar a empleados", totals.get("net", 0)),
        ("Aportes a cargo del empleador", totals.get("employer", 0)),
        ("Provisiones de prestaciones sociales", totals.get("provisions", 0)),
        ("Costo total del contrato fijo para el empleador", totals.get("employer_cost", 0)),
        ("Honorarios por servicio (liquidaciones aprobadas/pagadas)", totals.get("commissions", 0)),
        ("COSTO TOTAL DE PERSONAL", totals.get("total_personnel_cost", 0)),
    ]
    header_row(ws, 6, ["Concepto", "Valor"])
    row = 7
    for label, value in kpis:
        last = label.startswith("COSTO TOTAL DE")
        body_cell(ws, row, 1, label, strong=last, fill=soft_fill if last else None)
        body_cell(ws, row, 2, float(value or 0), money=True, strong=last, fill=soft_fill if last else None)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Desglose por concepto").font = Font(
        name="Calibri", size=11, bold=True, color=PRIMARY_DARK_HEX
    )
    row += 1
    header_row(ws, row, ["Concepto", "Valor"])
    row += 1
    groups = [
        ("Devengados", "earnings"),
        ("Deducciones del empleado", "employee_deductions"),
        ("Aportes del empleador", "employer_contributions"),
        ("Provisiones de prestaciones", "provisions"),
    ]
    for group_title, key in groups:
        body_cell(ws, row, 1, group_title, strong=True, fill=soft_fill)
        body_cell(ws, row, 2, None, fill=soft_fill)
        row += 1
        for (_, label), amount in sorted(_sum_codes(salaried, key).items(), key=lambda item: item[0][1]):
            body_cell(ws, row, 1, "   " + label)
            body_cell(ws, row, 2, round(amount, 2), money=True)
            row += 1
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 22
    ws.cell(row=row + 1, column=1, value=DISCLAIMER).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=row + 1, column=1).font = Font(name="Calibri", size=9, italic=True, color=MUTED_HEX)
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=2)
    ws.row_dimensions[row + 1].height = 48

    # ---------------- Detalle por empleado
    detail = wb.create_sheet("Detalle por empleado")
    columns = [
        "Empleado",
        "Documento",
        "Cargo",
        "Salario básico",
        "Días",
        "Devengado",
        "Salud empleado",
        "Pensión empleado",
        "Otras deducciones",
        "Neto a pagar",
        "Salud empleador",
        "Pensión empleador",
        "ARL",
        "Caja",
        "SENA + ICBF",
        "Cesantías",
        "Int. cesantías",
        "Prima",
        "Vacaciones",
        "Costo total empleador",
    ]
    title(detail, "Detalle por empleado", subtitle, len(columns))
    header_row(detail, 6, columns)
    r = 7
    for line in salaried:
        c = line["computed"]
        ded = {x["code"]: x["amount"] for x in c["employee_deductions"]}
        emp = {x["code"]: x["amount"] for x in c["employer_contributions"]}
        prov = {x["code"]: x["amount"] for x in c["provisions"]}
        other = c["deductions_total"] - ded.get("health", 0) - ded.get("pension", 0)
        values = [
            line.get("name"),
            line.get("document") or "",
            line.get("position") or "",
            c["base_salary"],
            c["days_worked"],
            c["gross"],
            ded.get("health", 0),
            ded.get("pension", 0),
            other,
            c["net_pay"],
            emp.get("health", 0),
            emp.get("pension", 0),
            emp.get("arl", 0),
            emp.get("ccf", 0),
            emp.get("sena", 0) + emp.get("icbf", 0),
            prov.get("cesantias", 0),
            prov.get("cesantias_interest", 0),
            prov.get("prima", 0),
            prov.get("vacations", 0),
            c["employer_cost"],
        ]
        for col, value in enumerate(values, start=1):
            body_cell(detail, r, col, value, money=col not in (1, 2, 3, 5))
        r += 1
    if salaried:
        body_cell(detail, r, 1, "TOTAL", strong=True, fill=soft_fill)
        for col in range(2, len(columns) + 1):
            if col in (2, 3, 5):
                body_cell(detail, r, col, None, fill=soft_fill)
            else:
                letter = get_column_letter(col)
                body_cell(detail, r, col, f"=SUM({letter}7:{letter}{r - 1})", money=True, strong=True, fill=soft_fill)
    else:
        detail.cell(row=7, column=1, value="Sin empleados con contrato fijo en este periodo.").font = base_font
    widths = [28, 14, 18] + [16] * (len(columns) - 3)
    for index, width in enumerate(widths, start=1):
        detail.column_dimensions[get_column_letter(index)].width = width
    detail.freeze_panes = "B7"
    detail.page_setup.orientation = "landscape"
    detail.page_setup.fitToWidth = 1
    detail.page_setup.fitToHeight = 0
    detail.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------- Honorarios por servicio
    if commission:
        com = wb.create_sheet("Contrato por servicio")
        title(com, "Contrato por servicio (comisión)", subtitle, 3)
        header_row(com, 6, ["Profesional", "Liquidaciones del periodo", "Valor"])
        r = 7
        for line in commission:
            body_cell(com, r, 1, line.get("name"))
            body_cell(com, r, 2, line.get("settlement_count", 0))
            body_cell(com, r, 3, float(line.get("settlement_total", 0)), money=True)
            r += 1
        com.column_dimensions["A"].width = 32
        com.column_dimensions["B"].width = 24
        com.column_dimensions["C"].width = 20

    # ---------------- Centros de costo (solo si se definieron)
    if any(x.get("cost_center") for x in run.get("lines", [])):
        cc = wb.create_sheet("Por centro de costos")
        title(cc, "Costo de personal por centro de costos", subtitle, 5)
        header_row(cc, 6, ["Centro de costos", "Empleados", "Devengado", "Aportes y provisiones", "Costo total"])
        groups: dict = {}
        for line in run.get("lines", []):
            key = line.get("cost_center") or "Sin centro de costos"
            bucket = groups.setdefault(key, {"count": 0, "gross": 0.0, "extra": 0.0, "cost": 0.0})
            bucket["count"] += 1
            computed = line.get("computed")
            if computed:
                bucket["gross"] += computed["gross"]
                bucket["extra"] += computed["employer_total"] + computed["provisions_total"]
                bucket["cost"] += computed["employer_cost"]
            else:
                bucket["gross"] += float(line.get("settlement_total", 0))
                bucket["cost"] += float(line.get("settlement_total", 0))
        r = 7
        for name, bucket in sorted(groups.items()):
            for col, value in enumerate(
                [name, bucket["count"], bucket["gross"], bucket["extra"], bucket["cost"]], start=1
            ):
                body_cell(cc, r, col, value, money=col >= 3)
            r += 1
        for letter, width in zip("ABCDE", (34, 12, 18, 22, 18)):
            cc.column_dimensions[letter].width = width

    # ---------------- Parametros y avisos
    info = wb.create_sheet("Parámetros y avisos")
    title(info, "Parámetros usados y avisos", subtitle, 2)
    params = run.get("params", {})
    rows = [
        ("Año de parámetros", params.get("year")),
        ("Salario mínimo (SMMLV)", fmt(params.get("smmlv"))),
        ("Auxilio de transporte", fmt(params.get("transport_aid"))),
        ("Empleador exonerado (art. 114-1 E.T.)", "Sí" if run.get("exonerated") else "No"),
        ("Frecuencia de pago", FREQUENCY_ES.get(run.get("frequency"), run.get("frequency"))),
    ]
    header_row(info, 6, ["Parámetro", "Valor"])
    r = 7
    for label, value in rows:
        body_cell(info, r, 1, label)
        body_cell(info, r, 2, value)
        r += 1
    r += 1
    for line in salaried:
        for note in line["computed"].get("notes", []):
            body_cell(info, r, 1, line.get("name"), strong=True)
            body_cell(info, r, 2, note)
            r += 1
    if run.get("corrections"):
        r += 1
        header_row(info, r, ["Corrección", "Motivo"])
        r += 1
        for item in run["corrections"]:
            body_cell(
                info,
                r,
                1,
                f"{item.get('at', '')[:10]} · de {STATUS_ES.get(item.get('from_status'), item.get('from_status'))}",
            )
            body_cell(info, r, 2, item.get("reason"))
            r += 1
    r += 1
    info.cell(row=r, column=1, value=DISCLAIMER).alignment = Alignment(wrap_text=True, vertical="top")
    info.cell(row=r, column=1).font = Font(name="Calibri", size=9, italic=True, color=MUTED_HEX)
    info.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    info.row_dimensions[r].height = 48
    info.column_dimensions["A"].width = 42
    info.column_dimensions["B"].width = 90

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


# ----------------------------------------------------------------------------------------------- PDF

PRIMARY = colors.HexColor("#" + PRIMARY_HEX)
PRIMARY_DARK = colors.HexColor("#" + PRIMARY_DARK_HEX)
SOFT = colors.HexColor("#" + SOFT_HEX)
BORDER = colors.HexColor("#" + BORDER_HEX)
TEXT = colors.HexColor("#" + TEXT_HEX)
MUTED = colors.HexColor("#" + MUTED_HEX)

_S = {
    "small": ParagraphStyle("small", fontName="Helvetica", fontSize=7.5, leading=10, textColor=MUTED),
    "label": ParagraphStyle("label", fontName="Helvetica-Bold", fontSize=7, leading=9, textColor=MUTED),
    "value": ParagraphStyle("value", fontName="Helvetica", fontSize=9, leading=12, textColor=TEXT),
    "h": ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=10, leading=13, textColor=PRIMARY_DARK),
}


def build_slip_pdf(run: dict, line: dict, organization: dict, logo_bytes: Optional[bytes] = None) -> bytes:
    """Colilla de pago (fondo blanco, colores de Nexus). Soporta contrato fijo y contrato por servicio."""
    out = BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
        title=f"Colilla de pago - {line.get('name')}",
        author=organization.get("name") or "Nexus",
    )
    width = A4[0] - 36 * mm
    story: list = []

    # Encabezado: logo definido por el negocio (o nombre) + titulo
    logo_cell = Paragraph(
        f"<b>{organization.get('name') or 'Nexus'}</b>",
        ParagraphStyle("o", fontName="Helvetica-Bold", fontSize=14, textColor=PRIMARY_DARK),
    )
    if logo_bytes:
        try:
            reader = ImageReader(BytesIO(logo_bytes))
            iw, ih = reader.getSize()
            height = 16 * mm
            from reportlab.platypus import Image

            logo_cell = Image(BytesIO(logo_bytes), width=height * iw / ih, height=height)
        except Exception:  # un logo ilegible nunca debe impedir generar la colilla
            pass
    head = Table(
        [
            [
                logo_cell,
                Paragraph(
                    "<b>COLILLA DE PAGO</b><br/>" + period_label(run),
                    ParagraphStyle("t", fontName="Helvetica", fontSize=10, leading=14, alignment=2, textColor=TEXT),
                ),
            ]
        ],
        colWidths=[width * 0.55, width * 0.45],
    )
    head.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LINEBELOW", (0, 0), (-1, 0), 1.2, PRIMARY),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story += [head, Spacer(1, 6 * mm)]

    contract = (
        "Contrato fijo (salario)" if line.get("contract_type") == "fixed_salary" else "Contrato por servicio (comisión)"
    )
    c = line.get("computed")
    info = [
        [Paragraph("EMPLEADO", _S["label"]), Paragraph("DOCUMENTO", _S["label"]), Paragraph("CARGO", _S["label"])],
        [
            Paragraph(line.get("name") or "", _S["value"]),
            Paragraph(line.get("document") or "—", _S["value"]),
            Paragraph(line.get("position") or "—", _S["value"]),
        ],
        [
            Paragraph("TIPO DE CONTRATO", _S["label"]),
            Paragraph("FRECUENCIA", _S["label"]),
            Paragraph("DÍAS PAGADOS" if c else "LIQUIDACIONES", _S["label"]),
        ],
        [
            Paragraph(contract, _S["value"]),
            Paragraph(FREQUENCY_ES.get(run.get("frequency"), ""), _S["value"]),
            Paragraph(f"{c['days_worked']:g}" if c else str(line.get("settlement_count", 0)), _S["value"]),
        ],
    ]
    t = Table(info, colWidths=[width * 0.4, width * 0.3, width * 0.3])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), SOFT),
                ("BOX", (0, 0), (-1, -1), 0.6, BORDER),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story += [t, Spacer(1, 5 * mm)]

    if c:
        earn = [[Paragraph("<b>DEVENGADOS</b>", _S["h"]), ""]] + [
            [Paragraph(r["label"], _S["value"]), fmt(r["amount"])] for r in c["earnings"]
        ]
        earn.append([Paragraph("<b>Total devengado</b>", _S["value"]), fmt(c["gross"])])
        ded = [[Paragraph("<b>DEDUCCIONES</b>", _S["h"]), ""]] + [
            [Paragraph(r["label"], _S["value"]), fmt(r["amount"])] for r in c["employee_deductions"]
        ]
        ded.append([Paragraph("<b>Total deducciones</b>", _S["value"]), fmt(c["deductions_total"])])
        half = (width - 6 * mm) / 2

        def block(rows):
            tb = Table(rows, colWidths=[half * 0.66, half * 0.34])
            tb.setStyle(
                TableStyle(
                    [
                        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                        ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
                        ("FONTSIZE", (1, 0), (1, -1), 9),
                        ("TEXTCOLOR", (1, 0), (1, -1), TEXT),
                        ("LINEBELOW", (0, 0), (-1, 0), 0.8, PRIMARY),
                        ("LINEABOVE", (0, -1), (-1, -1), 0.6, BORDER),
                        ("FONTNAME", (1, -1), (1, -1), "Helvetica-Bold"),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                )
            )
            return tb

        two = Table([[block(earn), block(ded)]], colWidths=[half + 3 * mm, half + 3 * mm])
        two.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3 * mm),
                ]
            )
        )
        story += [two, Spacer(1, 4 * mm)]
        net = Table(
            [
                [
                    Paragraph(
                        "<b>NETO A PAGAR</b>",
                        ParagraphStyle("n", fontName="Helvetica-Bold", fontSize=11, textColor=colors.white),
                    ),
                    Paragraph(
                        f"<b>{fmt(c['net_pay'])}</b>",
                        ParagraphStyle(
                            "nv", fontName="Helvetica-Bold", fontSize=13, textColor=colors.white, alignment=2
                        ),
                    ),
                ]
            ],
            colWidths=[width * 0.5, width * 0.5],
        )
        net.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), PRIMARY),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story += [net, Spacer(1, 5 * mm)]

        story.append(
            Paragraph("Aportes a cargo del empleador y prestaciones causadas (informativo, no se descuentan)", _S["h"])
        )
        rows = [[Paragraph(r["label"], _S["small"]), fmt(r["amount"])] for r in c["employer_contributions"]]
        rows += [[Paragraph(r["label"], _S["small"]), fmt(r["amount"])] for r in c["provisions"]]
        info_t = Table(rows, colWidths=[width * 0.7, width * 0.3])
        info_t.setStyle(
            TableStyle(
                [
                    ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                    ("FONTSIZE", (1, 0), (1, -1), 8),
                    ("TEXTCOLOR", (1, 0), (1, -1), MUTED),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.3, BORDER),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )
        story += [Spacer(1, 2 * mm), info_t]
        for note in c.get("notes", []):
            story += [Spacer(1, 2 * mm), Paragraph("• " + note, _S["small"])]
    else:
        total = float(line.get("settlement_total", 0))
        rows = [
            [Paragraph("<b>HONORARIOS POR SERVICIO</b>", _S["h"]), ""],
            [Paragraph("Liquidaciones aprobadas o pagadas en el periodo", _S["value"]), fmt(total)],
        ]
        tb = Table(rows, colWidths=[width * 0.7, width * 0.3])
        tb.setStyle(
            TableStyle(
                [
                    ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                    ("LINEBELOW", (0, 0), (-1, 0), 0.8, PRIMARY),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story += [
            tb,
            Spacer(1, 3 * mm),
            Paragraph(
                "Este contrato funciona por porcentaje de los servicios; el detalle está en las liquidaciones.",
                _S["small"],
            ),
        ]

    story += [Spacer(1, 8 * mm), Paragraph(DISCLAIMER, _S["small"])]

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setStrokeColor(BORDER)
        canvas.line(18 * mm, 12 * mm, A4[0] - 18 * mm, 12 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(MUTED)
        canvas.drawString(18 * mm, 8 * mm, f"{organization.get('name') or ''} · generado con Nexus by CS2")
        canvas.drawRightString(
            A4[0] - 18 * mm, 8 * mm, f"Versión {run.get('version', 1)} · {STATUS_ES.get(run.get('status'), '')}"
        )
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
