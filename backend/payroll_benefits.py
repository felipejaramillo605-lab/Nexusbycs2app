"""Prestaciones sociales, sabana de nomina y archivo de dispersion bancaria (consulta interna).

Los valores se toman de las nominas aprobadas o pagadas de esta app (provisiones causadas). El calculo legal de la
prima, las cesantias y sus intereses usa el salario promedio y puede diferir. NO es soporte de nomina electronica,
facturas electronicas ni de auditorias de la UGPP, y el archivo de dispersion es un CSV generico: no es un formato
bancario oficial.
"""

from __future__ import annotations

import csv
import uuid
from datetime import datetime, timezone
from io import BytesIO, StringIO
from typing import Dict, List, Literal, Optional

from fastapi import APIRouter, Cookie, Header, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pydantic import BaseModel

from payroll_co import DISCLAIMER
from payroll_reports import COP, MONTHS_ES, safe_text
from payroll_reports import PRIMARY_HEX as XLSX_HEAD
from payroll_reports import SOFT_HEX as XLSX_SOFT

KINDS = {
    "prima": {"label": "Prima de servicios", "provision": "prima", "semestral": True},
    "cesantias": {"label": "Cesantías", "provision": "cesantias", "semestral": False},
    "cesantias_interest": {"label": "Intereses a las cesantías", "provision": "cesantias_interest", "semestral": False},
}
NOTE = (
    "Valores causados según las nóminas aprobadas o pagadas en esta app. El cálculo legal usa el salario promedio del "
    "periodo y puede diferir: confírmalo con tu contador. " + DISCLAIMER
)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def deadline_text(kind: str, year: int, semester: Optional[int]) -> str:
    if kind == "prima":
        return f"Pagar a más tardar el {'30 de junio' if semester == 1 else '20 de diciembre'} de {year}"
    if kind == "cesantias":
        return f"Consignar al fondo a más tardar el 14 de febrero de {year + 1}"
    return f"Pagar al trabajador a más tardar el 31 de enero de {year + 1}"


def months_in_scope(kind: str, semester: Optional[int]) -> range:
    if KINDS[kind]["semestral"]:
        return range(1, 7) if semester == 1 else range(7, 13)
    return range(1, 13)


def aggregate_benefit(runs: List[dict], kind: str, year: int, semester: Optional[int]) -> List[dict]:
    """Suma, por empleado, la provision causada en las nominas aprobadas o pagadas del periodo."""
    code = KINDS[kind]["provision"]
    months = months_in_scope(kind, semester)
    rows: Dict[str, dict] = {}
    for run in runs:
        if run["status"] not in ("approved", "paid") or run["year"] != year or run["month"] not in months:
            continue
        for line in run["lines"]:
            computed = line.get("computed")
            if not computed:
                continue
            amount = next((p["amount"] for p in computed["provisions"] if p["code"] == code), 0.0)
            row = rows.setdefault(
                line["barber_id"],
                {
                    "barber_id": line["barber_id"],
                    "name": line["name"],
                    "document": line.get("document"),
                    "amount": 0.0,
                    "months": set(),
                },
            )
            row["amount"] += float(amount)
            row["months"].add(run["month"])
    result = []
    for row in rows.values():
        result.append({**row, "amount": round(row["amount"], 2), "months": sorted(row["months"])})
    return sorted(result, key=lambda r: r["name"] or "")


def build_benefit_workbook(org_name: str, kind: str, year: int, semester: Optional[int], rows: List[dict]) -> bytes:
    thin = Side(style="thin", color="E4E7EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    ws = wb.active
    ws.title = KINDS[kind]["label"][:30]
    ws.sheet_view.showGridLines = False
    ws["A1"] = org_name
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="6D28D9")
    ws["A2"] = f"{KINDS[kind]['label']} — {year}" + (f" (semestre {semester})" if semester else "")
    ws["A2"].font = Font(name="Calibri", size=12, bold=True)
    ws["A3"] = deadline_text(kind, year, semester)
    ws["A3"].font = Font(name="Calibri", size=10, color="667085")
    for index, label in enumerate(["Empleado", "Documento", "Meses incluidos", "Valor causado"], start=1):
        cell = ws.cell(row=5, column=index, value=label)
        cell.fill, cell.font, cell.border = (
            PatternFill("solid", fgColor=XLSX_HEAD),
            Font(name="Calibri", bold=True, color="FFFFFF"),
            border,
        )
    for offset, row in enumerate(rows, start=6):
        values = [
            row["name"],
            row.get("document") or "",
            ", ".join(MONTHS_ES[m - 1] for m in row["months"]),
            row["amount"],
        ]
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row=offset, column=index, value=safe_text(value))
            cell.border = border
            if index == 4:
                cell.number_format = COP
    total_row = 6 + len(rows)
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(name="Calibri", bold=True)
    total = ws.cell(row=total_row, column=4, value=f"=SUM(D6:D{total_row - 1})" if rows else 0)
    total.number_format, total.font = COP, Font(name="Calibri", bold=True)
    for col in range(1, 5):
        ws.cell(row=total_row, column=col).fill = PatternFill("solid", fgColor=XLSX_SOFT)
    note = ws.cell(row=total_row + 2, column=1, value=NOTE)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    note.font = Font(name="Calibri", size=9, italic=True, color="667085")
    ws.merge_cells(start_row=total_row + 2, start_column=1, end_row=total_row + 2, end_column=4)
    ws.row_dimensions[total_row + 2].height = 60
    for letter, width in zip("ABCD", (32, 18, 38, 20)):
        ws.column_dimensions[letter].width = width
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def build_sabana_workbook(org_name: str, year: int, runs: List[dict]) -> bytes:
    """Una fila por empleado y mes con devengado, deducciones, neto, aportes del empleador y provisiones."""
    thin = Side(style="thin", color="E4E7EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sábana de nómina"
    ws.sheet_view.showGridLines = False
    ws["A1"] = org_name
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="6D28D9")
    ws["A2"] = f"Sábana de nómina {year} (consulta interna)"
    ws["A2"].font = Font(name="Calibri", size=12, bold=True)
    ws["A3"] = "Esta app no sirve como soporte de nómina electrónica, facturas electrónicas ni para la UGPP."
    ws["A3"].font = Font(name="Calibri", size=10, bold=True, color="B54708")
    headers = [
        "Empleado",
        "Documento",
        "Mes",
        "Periodo",
        "Estado",
        "Devengado",
        "Deducciones",
        "Neto",
        "Aportes empleador",
        "Provisiones",
        "Costo empleador",
    ]
    for index, label in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=index, value=label)
        cell.fill, cell.font, cell.border = (
            PatternFill("solid", fgColor=XLSX_HEAD),
            Font(name="Calibri", bold=True, color="FFFFFF"),
            border,
        )
    row_index = 6
    for run in sorted(runs, key=lambda r: (r["month"], r.get("half") or 0)):
        if run["year"] != year or run["status"] not in ("approved", "paid"):
            continue
        for line in sorted(run["lines"], key=lambda x: x.get("name") or ""):
            computed = line.get("computed")
            if not computed:
                continue
            values = [
                line["name"],
                line.get("document") or "",
                MONTHS_ES[run["month"] - 1].capitalize(),
                run["number"],
                run["status"],
                computed["gross"],
                computed["deductions_total"],
                computed["net_pay"],
                computed["employer_total"],
                computed["provisions_total"],
                computed["employer_cost"],
            ]
            for index, value in enumerate(values, start=1):
                cell = ws.cell(row=row_index, column=index, value=safe_text(value))
                cell.border = border
                if index >= 6:
                    cell.number_format = COP
            row_index += 1
    for index, width in enumerate((30, 16, 14, 18, 12, 16, 16, 16, 18, 16, 18), start=1):
        ws.column_dimensions[chr(64 + index)].width = width
    ws.freeze_panes = "A6"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def dispersion_rows(run: dict, contracts: Dict[str, dict]) -> dict:
    """Pagos netos listos para el banco y quienes no tienen datos bancarios completos."""
    ready, missing = [], []
    for line in run["lines"]:
        computed = line.get("computed")
        if not computed or computed["net_pay"] <= 0:
            continue
        contract = contracts.get(line["barber_id"]) or {}
        data = {
            k: (contract.get(k) or "").strip() if isinstance(contract.get(k), str) else contract.get(k)
            for k in ("bank_name", "account_type", "account_number")
        }
        if not (data["bank_name"] and data["account_type"] and data["account_number"] and line.get("document")):
            missing.append(line["name"])
            continue
        ready.append(
            {
                "document": line["document"],
                "name": line["name"],
                **data,
                "amount": round(computed["net_pay"]),
                "reference": run["number"],
            }
        )
    return {"ready": ready, "missing": missing}


class ApplyIn(BaseModel):
    kind: Literal["prima", "cesantias_interest"]
    year: int
    semester: Optional[Literal[1, 2]] = None
    organization_id: Optional[str] = None


def build_benefits_router(db, get_current_user, require_management_role, resolve_team_organization):
    router = APIRouter()

    async def manager_org(user, requested):
        require_management_role(user)
        return await resolve_team_organization(user, requested)

    async def org_runs(org_id):
        return await db.payroll_runs.find({"organization_id": org_id}, {"_id": 0}).to_list(2000)

    async def org_name(org_id):
        organization = await db.organizations.find_one({"organization_id": org_id}, {"_id": 0})
        return (organization or {}).get("name") or "Nexus"

    def check(kind, semester):
        if kind not in KINDS:
            raise HTTPException(status_code=404, detail="Prestación no encontrada")
        if KINDS[kind]["semestral"] and semester not in (1, 2):
            raise HTTPException(status_code=400, detail="Indica el semestre (1 o 2)")

    @router.get("/payroll/benefits/{kind}", tags=["payroll"])
    async def preview(
        kind: str,
        year: int,
        semester: Optional[int] = None,
        organization_id: Optional[str] = None,
        authorization: Optional[str] = Header(None),
        session_token: Optional[str] = Cookie(None),
    ):
        user = await get_current_user(authorization, session_token)
        org_id = await manager_org(user, organization_id)
        check(kind, semester)
        rows = aggregate_benefit(await org_runs(org_id), kind, year, semester)
        return {
            "kind": kind,
            "label": KINDS[kind]["label"],
            "year": year,
            "semester": semester,
            "deadline": deadline_text(kind, year, semester),
            "items": rows,
            "total": round(sum(r["amount"] for r in rows), 2),
            "note": NOTE,
        }

    @router.get("/payroll/benefits/{kind}/export.xlsx", tags=["payroll"])
    async def export_benefit(
        kind: str,
        year: int,
        semester: Optional[int] = None,
        organization_id: Optional[str] = None,
        authorization: Optional[str] = Header(None),
        session_token: Optional[str] = Cookie(None),
    ):
        user = await get_current_user(authorization, session_token)
        org_id = await manager_org(user, organization_id)
        check(kind, semester)
        rows = aggregate_benefit(await org_runs(org_id), kind, year, semester)
        content = build_benefit_workbook(await org_name(org_id), kind, year, semester, rows)
        tag = f"{year}" + (f"-S{semester}" if semester else "")
        return Response(
            content,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{kind}_{tag}.xlsx"'},
        )

    @router.post("/payroll/benefits/apply", tags=["payroll"])
    async def apply_to_payroll(
        data: ApplyIn, authorization: Optional[str] = Header(None), session_token: Optional[str] = Cookie(None)
    ):
        """Registra el pago de prima o intereses como novedad aprobada (no constituye salario)."""
        user = await get_current_user(authorization, session_token)
        org_id = await manager_org(user, data.organization_id)
        check(data.kind, data.semester)
        rows = aggregate_benefit(await org_runs(org_id), data.kind, data.year, data.semester)
        tag = f"{data.year}" + (f" semestre {data.semester}" if data.semester else "")
        concept = f"{KINDS[data.kind]['label']} {tag}"
        existing = await db.payroll_novelties.find(
            {"organization_id": org_id, "type": "bonus", "concept": concept}, {"_id": 0}
        ).to_list(2000)
        done = {r["barber_id"] for r in existing if r["status"] in ("pending", "approved")}
        created = 0
        for row in rows:
            if row["barber_id"] in done or row["amount"] <= 0:
                continue
            await db.payroll_novelties.insert_one(
                {
                    "novelty_id": f"nov_{uuid.uuid4().hex[:14]}",
                    "organization_id": org_id,
                    "barber_id": row["barber_id"],
                    "employee_name": row["name"],
                    "type": "bonus",
                    "concept": concept,
                    "date": datetime.now(timezone.utc).date().isoformat(),
                    "amount": row["amount"],
                    "constitutes_salary": False,
                    "note": "Generado desde Prestaciones sociales",
                    "status": "approved",
                    "applied_run_id": None,
                    "decided_by": user.user_id,
                    "created_by": user.user_id,
                    "created_at": _now(),
                }
            )
            created += 1
        return {"created": created, "skipped": len(rows) - created}

    @router.get("/payroll/sabana.xlsx", tags=["payroll"])
    async def sabana(
        year: int,
        organization_id: Optional[str] = None,
        authorization: Optional[str] = Header(None),
        session_token: Optional[str] = Cookie(None),
    ):
        user = await get_current_user(authorization, session_token)
        org_id = await manager_org(user, organization_id)
        content = build_sabana_workbook(await org_name(org_id), year, await org_runs(org_id))
        return Response(
            content,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="sabana_nomina_{year}.xlsx"'},
        )

    async def dispersion_data(org_id, run_id):
        run = await db.payroll_runs.find_one({"run_id": run_id, "organization_id": org_id}, {"_id": 0})
        if not run:
            raise HTTPException(status_code=404, detail="Corrida de nómina no encontrada")
        contracts = {
            c["barber_id"]: c
            for c in await db.payroll_contracts.find({"organization_id": org_id}, {"_id": 0}).to_list(500)
        }
        return run, dispersion_rows(run, contracts)

    @router.get("/payroll/runs/{run_id}/dispersion", tags=["payroll"])
    async def dispersion_preview(
        run_id: str,
        organization_id: Optional[str] = None,
        authorization: Optional[str] = Header(None),
        session_token: Optional[str] = Cookie(None),
    ):
        user = await get_current_user(authorization, session_token)
        _, data = await dispersion_data(await manager_org(user, organization_id), run_id)
        return {
            **data,
            "total": sum(r["amount"] for r in data["ready"]),
            "note": "CSV genérico: no es un formato bancario oficial.",
        }

    @router.get("/payroll/runs/{run_id}/dispersion.csv", tags=["payroll"])
    async def dispersion_csv(
        run_id: str,
        organization_id: Optional[str] = None,
        authorization: Optional[str] = Header(None),
        session_token: Optional[str] = Cookie(None),
    ):
        user = await get_current_user(authorization, session_token)
        run, data = await dispersion_data(await manager_org(user, organization_id), run_id)
        if run["status"] not in ("approved", "paid"):
            raise HTTPException(status_code=409, detail="Aprueba la nómina antes de generar el archivo de pagos")
        buffer = StringIO()
        writer = csv.writer(buffer, delimiter=";")
        writer.writerow(["documento", "nombre", "banco", "tipo_cuenta", "numero_cuenta", "valor_neto", "referencia"])
        for row in data["ready"]:
            writer.writerow(
                [
                    safe_text(row["document"]),
                    safe_text(row["name"]),
                    safe_text(row["bank_name"]),
                    row["account_type"],
                    safe_text(row["account_number"]),
                    row["amount"],
                    row["reference"],
                ]
            )
        return Response(
            "﻿" + buffer.getvalue(),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="dispersion_{run["number"]}.csv"'},
        )

    return router
