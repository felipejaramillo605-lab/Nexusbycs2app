"""Modulo de nomina: contratos, auxilios, corridas, correcciones, Excel y colillas (base simulada)."""

import sys
from datetime import date
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import payroll as subject  # noqa: E402
import payroll_co  # noqa: E402

SMMLV = payroll_co.params_for(2026)["smmlv"]


def matches(doc, key, condition):
    value = doc.get(key)
    if isinstance(condition, dict):
        if "$ne" in condition:
            return value != condition["$ne"]
        return True
    return value == condition


class Cursor:
    def __init__(self, rows):
        self.rows = rows

    async def to_list(self, limit):
        return [dict(r) for r in self.rows[:limit]]


class Coll:
    def __init__(self, docs=()):
        self.docs = [dict(d) for d in docs]

    def _select(self, query):
        return [d for d in self.docs if all(matches(d, k, v) for k, v in query.items())]

    def find(self, query, projection=None):
        return Cursor(self._select(query))

    async def find_one(self, query, projection=None):
        found = self._select(query)
        return dict(found[0]) if found else None

    async def insert_one(self, doc):
        self.docs.append(dict(doc))

    async def update_one(self, query, update):
        for doc in self._select(query):
            doc.update(update.get("$set", {}))
            return

    async def delete_one(self, query):
        found = self._select(query)
        if found:
            self.docs.remove(found[0])


USERS = {
    "boss": SimpleNamespace(user_id="u_boss", role="manager", organization_id="org_a"),
    "ana": SimpleNamespace(user_id="u_ana", role="staff", organization_id="org_a"),
    "luis": SimpleNamespace(user_id="u_luis", role="staff", organization_id="org_a"),
}


def build():
    db = SimpleNamespace(
        barbers=Coll(
            [
                {"organization_id": "org_a", "barber_id": "b1", "name": "Ana", "user_id": "u_ana"},
                {"organization_id": "org_a", "barber_id": "b2", "name": "Luis", "user_id": "u_luis"},
                {"organization_id": "org_b", "barber_id": "b9", "name": "Otra org"},
            ]
        ),
        organizations=Coll([{"organization_id": "org_a", "name": "Barbería Norte"}]),
        payroll_settings=Coll(),
        payroll_extras=Coll(),
        payroll_contracts=Coll(),
        payroll_runs=Coll(),
        staff_settlements=Coll(
            [
                {
                    "organization_id": "org_a",
                    "barber_id": "b2",
                    "status": "paid",
                    "period_end": "2026-10-20",
                    "total_amount": 800000,
                },
                {
                    "organization_id": "org_a",
                    "barber_id": "b2",
                    "status": "draft",
                    "period_end": "2026-10-25",
                    "total_amount": 999999,
                },
                {
                    "organization_id": "org_a",
                    "barber_id": "b2",
                    "status": "paid",
                    "period_end": "2026-09-30",
                    "total_amount": 111111,
                },
            ]
        ),
    )

    async def get_current_user(authorization=None, session_token=None):
        user = USERS.get(authorization)
        if not user:
            raise HTTPException(status_code=401, detail="Not authenticated")
        return user

    def require_management_role(user):
        if user.role not in ("manager", "admin", "owner"):
            raise HTTPException(status_code=403, detail="Management only")

    async def resolve_team_organization(user, requested):
        return user.organization_id

    app = FastAPI()
    app.include_router(
        subject.build_payroll_router(db, get_current_user, require_management_role, resolve_team_organization),
        prefix="/api",
    )
    return db, TestClient(app)


def H(who="boss"):
    return {"Authorization": who}


def fixed_contract(client, barber="b1", salary=SMMLV, **extra):
    body = {
        "contract_type": "fixed_salary",
        "base_salary": salary,
        "document": "1.020.333.444",
        "position": "Barbera",
        **extra,
    }
    return client.put(f"/api/payroll/contracts/{barber}", json=body, headers=H())


def make_run(client, **extra):
    response = client.post("/api/payroll/runs", json={"year": 2026, "month": 10, **extra}, headers=H())
    assert response.status_code == 200, response.text
    return response.json()


# ---------------------------------------------------------------- logica pura


def test_days_in_period_prorates_new_hires_for_month_and_half_month():
    start, end = date(2026, 10, 1), date(2026, 10, 31)
    assert subject.days_in_period(start, end, None, "monthly") == 30
    assert subject.days_in_period(start, end, date(2026, 9, 1), "monthly") == 30
    assert subject.days_in_period(start, end, date(2026, 10, 11), "monthly") == 20
    assert subject.days_in_period(start, end, date(2026, 11, 2), "monthly") == 0
    assert subject.days_in_period(date(2026, 10, 16), date(2026, 10, 31), date(2026, 10, 21), "biweekly") == 10
    assert subject.days_in_period(date(2026, 10, 1), date(2026, 10, 15), date(2026, 10, 1), "biweekly") == 15


# ---------------------------------------------------------------- ajustes, extras, contratos


def test_settings_default_and_overrides_and_only_management_can_use_them():
    _, client = build()
    body = client.get("/api/payroll/settings", headers=H()).json()
    assert (
        body["exonerated"] is True
        and body["params"]["smmlv"] > 0
        and "No reemplaza" in body["disclaimer"]
        or "no reemplaza" in body["disclaimer"].lower()
    )
    saved = client.put(
        "/api/payroll/settings",
        json={
            "exonerated": False,
            "default_arl_class": "III",
            "params_overrides": {"2026": {"smmlv": 2000000, "transport_aid": 260000}},
        },
        headers=H(),
    )
    assert saved.status_code == 200 and client.get("/api/payroll/settings", headers=H()).json()["exonerated"] is False
    assert (
        client.put("/api/payroll/settings", json={"params_overrides": {"2026": {"smmlv": 0}}}, headers=H()).status_code
        == 400
    )
    assert client.get("/api/payroll/settings", headers=H("ana")).status_code == 403


def test_extras_can_be_created_changed_and_deleted_with_validation():
    _, client = build()
    created = client.post(
        "/api/payroll/extras", json={"name": "Auxilio de internet", "kind": "fixed", "value": 50000}, headers=H()
    )
    assert created.status_code == 200
    extra_id = created.json()["extra_id"]
    assert (
        client.post(
            "/api/payroll/extras", json={"name": "Rodamiento", "kind": "percent", "value": 150}, headers=H()
        ).status_code
        == 400
    )
    assert (
        client.post(
            "/api/payroll/extras",
            json={"name": "Celular", "kind": "fixed", "value": 30000, "applies_to": "selected"},
            headers=H(),
        ).status_code
        == 400
    )
    changed = client.put(
        f"/api/payroll/extras/{extra_id}", json={"name": "Internet y datos", "kind": "percent", "value": 5}, headers=H()
    )
    assert changed.json()["name"] == "Internet y datos" and changed.json()["kind"] == "percent"
    assert len(client.get("/api/payroll/extras", headers=H()).json()["items"]) == 1
    assert client.delete(f"/api/payroll/extras/{extra_id}", headers=H()).json() == {"deleted": True}
    assert client.delete(f"/api/payroll/extras/{extra_id}", headers=H()).status_code == 404


def test_fixed_contract_must_respect_the_minimum_wage_and_listing_shows_the_type():
    _, client = build()
    assert fixed_contract(client, salary=SMMLV - 1000).status_code == 400
    assert (
        client.put("/api/payroll/contracts/b1", json={"contract_type": "fixed_salary"}, headers=H()).status_code == 400
    )
    assert fixed_contract(client).status_code == 200
    assert (
        client.put("/api/payroll/contracts/zzz", json={"contract_type": "service_commission"}, headers=H()).status_code
        == 404
    )
    rows = {r["barber_id"]: r for r in client.get("/api/payroll/contracts", headers=H()).json()["items"]}
    assert rows["b1"]["contract_type"] == "fixed_salary" and rows["b1"]["configured"] is True
    assert rows["b2"]["contract_type"] == "service_commission" and rows["b2"]["configured"] is False
    assert "b9" not in rows


# ---------------------------------------------------------------- corridas


def test_run_mixes_a_fixed_salary_line_and_commission_staff_and_blocks_duplicates():
    _, client = build()
    fixed_contract(client)
    client.put("/api/payroll/contracts/b2", json={"contract_type": "service_commission"}, headers=H())
    client.post(
        "/api/payroll/extras", json={"name": "Auxilio de internet", "kind": "fixed", "value": 50000}, headers=H()
    )
    run = make_run(client)
    lines = {x["barber_id"]: x for x in run["lines"]}
    assert lines["b1"]["computed"]["gross"] == round(SMMLV + payroll_co.params_for(2026)["transport_aid"] + 50000, 2)
    assert (
        lines["b2"]["settlement_total"] == 800000 and lines["b2"]["settlement_count"] == 1
    )  # solo la pagada de octubre
    assert run["totals"]["commissions"] == 800000
    assert run["totals"]["total_personnel_cost"] == round(run["totals"]["employer_cost"] + 800000, 2)
    assert run["number"] == "NOM-2026-10" and run["status"] == "draft"
    again = client.post("/api/payroll/runs", json={"year": 2026, "month": 10}, headers=H())
    assert again.status_code == 409


def test_biweekly_run_uses_only_biweekly_contracts_and_needs_the_half():
    _, client = build()
    fixed_contract(client, pay_frequency="biweekly")
    assert (
        client.post(
            "/api/payroll/runs", json={"year": 2026, "month": 10, "frequency": "biweekly"}, headers=H()
        ).status_code
        == 400
    )
    run = make_run(client, frequency="biweekly", half=1)
    assert run["number"] == "NOM-2026-10-Q1" and run["lines"][0]["computed"]["days_worked"] == 15
    assert run["lines"][0]["computed"]["gross"] == round((SMMLV + payroll_co.params_for(2026)["transport_aid"]) / 2, 2)
    assert all(x["computed"] is None for x in make_run(client, frequency="monthly")["lines"])  # contrato quincenal


def test_draft_lines_can_be_corrected_and_approved_runs_are_locked_until_reopened_with_a_reason():
    _, client = build()
    fixed_contract(client)
    run = make_run(client)
    rid = run["run_id"]
    edit = client.put(
        f"/api/payroll/runs/{rid}/lines/b1",
        json={"days_worked": 20, "adjustments": [{"label": "Préstamo", "amount": 50000, "kind": "deduction"}]},
        headers=H(),
    )
    assert edit.status_code == 200
    line = edit.json()["lines"][0]["computed"]
    assert line["days_worked"] == 20 and line["deductions_total"] == round(
        sum(x["amount"] for x in line["employee_deductions"]), 2
    )
    assert client.put(f"/api/payroll/runs/{rid}/lines/b1", json={"days_worked": 31}, headers=H()).status_code == 422
    assert client.put(f"/api/payroll/runs/{rid}/lines/b2", json={"days_worked": 10}, headers=H()).status_code == 404

    assert client.post(f"/api/payroll/runs/{rid}/pay", headers=H()).status_code == 409  # primero se aprueba
    assert client.post(f"/api/payroll/runs/{rid}/approve", headers=H()).json()["status"] == "approved"
    assert client.put(f"/api/payroll/runs/{rid}/lines/b1", json={"days_worked": 25}, headers=H()).status_code == 409
    assert client.post(f"/api/payroll/runs/{rid}/reopen", json={"reason": "no"}, headers=H()).status_code == 422
    reopened = client.post(
        f"/api/payroll/runs/{rid}/reopen", json={"reason": "Faltó registrar un descuento"}, headers=H()
    ).json()
    assert reopened["status"] == "draft" and reopened["version"] == 2
    assert (
        reopened["corrections"][0]["reason"] == "Faltó registrar un descuento"
        and reopened["corrections"][0]["from_status"] == "approved"
    )
    assert client.put(f"/api/payroll/runs/{rid}/lines/b1", json={"days_worked": 25}, headers=H()).status_code == 200
    client.post(f"/api/payroll/runs/{rid}/approve", headers=H())
    assert client.post(f"/api/payroll/runs/{rid}/pay", headers=H()).json()["status"] == "paid"
    assert client.post(f"/api/payroll/runs/{rid}/cancel", headers=H()).status_code == 409  # solo borradores


# ---------------------------------------------------------------- Excel y colillas


def test_excel_report_has_the_report_sheets_totals_and_corrections():
    _, client = build()
    fixed_contract(client)
    run = make_run(client)
    rid = run["run_id"]
    client.post(f"/api/payroll/runs/{rid}/approve", headers=H())
    client.post(f"/api/payroll/runs/{rid}/reopen", json={"reason": "Ajuste de días"}, headers=H())
    response = client.get(f"/api/payroll/runs/{rid}/report.xlsx", headers=H())
    assert response.status_code == 200 and "spreadsheetml" in response.headers["content-type"]
    assert "NOM-2026-10_gastos_de_personal.xlsx" in response.headers["content-disposition"]
    wb = load_workbook(BytesIO(response.content))
    assert wb.sheetnames[:2] == ["Resumen", "Detalle por empleado"]
    assert "Parámetros y avisos" in wb.sheetnames
    summary = " ".join(str(c.value) for row in wb["Resumen"].iter_rows() for c in row if c.value)
    assert "Barbería Norte" in summary and "COSTO TOTAL DE PERSONAL" in summary and "Octubre de 2026" in summary
    detail = wb["Detalle por empleado"]
    assert (
        detail["A7"].value == "Ana"
        and str(detail["R8"].value).startswith("=SUM")
        or str(detail["T8"].value).startswith("=SUM")
    )
    info = " ".join(str(c.value) for row in wb["Parámetros y avisos"].iter_rows() for c in row if c.value)
    assert "Ajuste de días" in info and "no reemplaza" in info.lower()


def test_manager_slip_pdf_and_staff_only_see_their_own_approved_slips():
    _, client = build()
    fixed_contract(client)
    client.put("/api/payroll/contracts/b2", json={"contract_type": "service_commission"}, headers=H())
    run = make_run(client)
    rid = run["run_id"]
    pdf = client.get(f"/api/payroll/runs/{rid}/slips/b1.pdf", headers=H())
    assert (
        pdf.status_code == 200
        and pdf.content.startswith(b"%PDF")
        and "colilla_NOM-2026-10_Ana.pdf" in pdf.headers["content-disposition"]
    )
    assert client.get(f"/api/payroll/runs/{rid}/slips/b2.pdf", headers=H()).content.startswith(
        b"%PDF"
    )  # contrato por servicio
    assert client.get(f"/api/payroll/runs/{rid}/slips/nadie.pdf", headers=H()).status_code == 404
    assert client.get(f"/api/payroll/runs/{rid}/slips/b1.pdf", headers=H("ana")).status_code == 403

    assert client.get("/api/staff/payroll/slips", headers=H("ana")).json()["items"] == []  # borrador: no visible
    assert client.get(f"/api/staff/payroll/slips/{rid}.pdf", headers=H("ana")).status_code == 404
    client.post(f"/api/payroll/runs/{rid}/approve", headers=H())
    mine = client.get("/api/staff/payroll/slips", headers=H("ana")).json()["items"]
    assert len(mine) == 1 and mine[0]["net_pay"] > 0 and mine[0]["label"] == "Octubre de 2026"
    assert client.get(f"/api/staff/payroll/slips/{rid}.pdf", headers=H("ana")).content.startswith(b"%PDF")
    commission = client.get("/api/staff/payroll/slips", headers=H("luis")).json()["items"]
    assert commission[0]["commission_total"] == 800000 and commission[0]["contract_type"] == "service_commission"
