"""Prestaciones sociales, sabana de nomina y dispersion (base simulada)."""

import sys
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import payroll_benefits as subject  # noqa: E402
import payroll_co  # noqa: E402

P = payroll_co.params_for(2026)


def computed(salary):
    return payroll_co.compute_line(base_salary=salary, frequency="monthly", risk_class="I", params=P)


def run(month, status="approved", year=2026, with_luis=True):
    lines = [
        {
            "barber_id": "b1",
            "name": "Ana",
            "document": "123",
            "contract_type": "fixed_salary",
            "computed": computed(P["smmlv"]),
        }
    ]
    if with_luis:
        lines.append(
            {
                "barber_id": "b2",
                "name": "Luis",
                "document": None,
                "contract_type": "fixed_salary",
                "computed": computed(2_500_000),
            }
        )
    lines.append(
        {
            "barber_id": "b3",
            "name": "Por servicio",
            "contract_type": "service_commission",
            "computed": None,
            "settlement_total": 500000,
        }
    )
    return {
        "run_id": f"pay_{year}_{month}",
        "organization_id": "org_a",
        "number": f"NOM-{year}-{month:02d}",
        "year": year,
        "month": month,
        "frequency": "monthly",
        "half": None,
        "status": status,
        "version": 1,
        "lines": lines,
    }


class Cursor:
    def __init__(self, rows):
        self.rows = rows

    async def to_list(self, limit):
        return [dict(r) for r in self.rows[:limit]]


class Coll:
    def __init__(self, docs=()):
        self.docs = [dict(d) for d in docs]

    def _select(self, query):
        return [d for d in self.docs if all(d.get(k) == v for k, v in query.items())]

    def find(self, query, projection=None):
        return Cursor(self._select(query))

    async def find_one(self, query, projection=None):
        found = self._select(query)
        return dict(found[0]) if found else None

    async def insert_one(self, doc):
        self.docs.append(dict(doc))


def build(runs):
    db = SimpleNamespace(
        payroll_runs=Coll(runs),
        payroll_novelties=Coll(),
        organizations=Coll([{"organization_id": "org_a", "name": "Barbería Norte"}]),
        payroll_contracts=Coll(
            [
                {
                    "organization_id": "org_a",
                    "barber_id": "b1",
                    "bank_name": "Bancolombia",
                    "account_type": "savings",
                    "account_number": "123456789",
                },
                {"organization_id": "org_a", "barber_id": "b2", "bank_name": "Davivienda"},
            ]
        ),
    )
    users = {
        "boss": SimpleNamespace(user_id="u1", role="manager", organization_id="org_a"),
        "ana": SimpleNamespace(user_id="u2", role="staff", organization_id="org_a"),
    }

    async def get_current_user(authorization=None, session_token=None):
        return users[authorization or "boss"]

    def require_management_role(user):
        if user.role != "manager":
            raise HTTPException(status_code=403, detail="Management only")

    async def resolve_team_organization(user, requested):
        return user.organization_id

    app = FastAPI()
    app.include_router(
        subject.build_benefits_router(db, get_current_user, require_management_role, resolve_team_organization),
        prefix="/api",
    )
    return db, TestClient(app)


def H(who="boss"):
    return {"Authorization": who}


def test_prima_adds_only_the_semester_months_of_approved_or_paid_runs():
    runs = [run(m) for m in (1, 2)] + [run(7), run(8, "paid"), run(9, "draft"), run(10, year=2025)]
    rows = subject.aggregate_benefit(runs, "prima", 2026, 2)
    ana = next(r for r in rows if r["barber_id"] == "b1")
    monthly = next(p["amount"] for p in computed(P["smmlv"])["provisions"] if p["code"] == "prima")
    assert ana["months"] == [7, 8] and ana["amount"] == round(monthly * 2, 2)
    assert all(r["barber_id"] != "b3" for r in rows)  # el contrato por servicio no causa prestaciones


def test_cesantias_use_the_whole_year_and_deadlines_are_stated():
    rows = subject.aggregate_benefit([run(m) for m in range(1, 13)], "cesantias", 2026, None)
    assert rows[0]["months"] == list(range(1, 13))
    assert "14 de febrero de 2027" in subject.deadline_text("cesantias", 2026, None)
    assert "31 de enero de 2027" in subject.deadline_text("cesantias_interest", 2026, None)
    assert "30 de junio de 2026" in subject.deadline_text(
        "prima", 2026, 1
    ) and "20 de diciembre" in subject.deadline_text("prima", 2026, 2)


def test_preview_and_export_with_validation_and_disclaimer():
    _, client = build([run(7), run(8)])
    body = client.get("/api/payroll/benefits/prima", params={"year": 2026, "semester": 2}, headers=H()).json()
    assert (
        body["label"] == "Prima de servicios" and body["total"] == round(sum(r["amount"] for r in body["items"]), 2) > 0
    )
    assert "UGPP" in body["note"] and "20 de diciembre" in body["deadline"]
    assert client.get("/api/payroll/benefits/prima", params={"year": 2026}, headers=H()).status_code == 400
    assert client.get("/api/payroll/benefits/otra", params={"year": 2026}, headers=H()).status_code == 404
    assert (
        client.get("/api/payroll/benefits/prima", params={"year": 2026, "semester": 2}, headers=H("ana")).status_code
        == 403
    )
    xlsx = client.get("/api/payroll/benefits/prima/export/xlsx", params={"year": 2026, "semester": 2}, headers=H())
    sheet = load_workbook(BytesIO(xlsx.content)).active
    assert sheet["A1"].value == "Barbería Norte" and sheet["A6"].value == "Ana" and "agosto" in sheet["C6"].value
    assert str(sheet["D8"].value).startswith("=SUM")


def test_applying_prima_creates_approved_non_salary_novelties_once():
    db, client = build([run(7), run(8)])
    first = client.post(
        "/api/payroll/benefits/apply", json={"kind": "prima", "year": 2026, "semester": 2}, headers=H()
    ).json()
    assert first == {"created": 2, "skipped": 0}
    novelty = db.payroll_novelties.docs[0]
    assert (
        novelty["status"] == "approved"
        and novelty["constitutes_salary"] is False
        and novelty["concept"] == "Prima de servicios 2026 semestre 2"
    )
    assert client.post(
        "/api/payroll/benefits/apply", json={"kind": "prima", "year": 2026, "semester": 2}, headers=H()
    ).json() == {"created": 0, "skipped": 2}
    assert (
        client.post("/api/payroll/benefits/apply", json={"kind": "cesantias", "year": 2026}, headers=H()).status_code
        == 422
    )  # van al fondo, no a la nómina


def test_sabana_lists_every_approved_month_per_employee_with_the_warning():
    _, client = build([run(1), run(2, "paid"), run(3, "draft")])
    sheet = load_workbook(
        BytesIO(client.get("/api/payroll/sabana/xlsx", params={"year": 2026}, headers=H()).content)
    ).active
    assert "ni para la UGPP" in sheet["A3"].value
    rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=6) if row[0].value]
    assert [r[0] for r in rows] == ["Ana", "Luis", "Ana", "Luis"] and {r[2] for r in rows} == {"Enero", "Febrero"}


def test_dispersion_lists_missing_bank_data_and_only_approved_runs_produce_the_file():
    runs = [run(10), run(11, "draft")]
    _, client = build(runs)
    preview = client.get("/api/payroll/runs/pay_2026_10/dispersion", headers=H()).json()
    assert [r["name"] for r in preview["ready"]] == ["Ana"] and preview["missing"] == ["Luis"]
    assert preview["total"] == preview["ready"][0]["amount"] and "no es un formato bancario" in preview["note"]
    file = client.get("/api/payroll/runs/pay_2026_10/dispersion/csv", headers=H())
    text = file.content.decode("utf-8-sig").splitlines()
    assert text[0] == "documento;nombre;banco;tipo_cuenta;numero_cuenta;valor_neto;referencia"
    assert text[1].startswith("123;Ana;Bancolombia;savings;123456789;") and len(text) == 2
    assert client.get("/api/payroll/runs/pay_2026_11/dispersion/csv", headers=H()).status_code == 409
    assert client.get("/api/payroll/runs/nada/dispersion", headers=H()).status_code == 404


def test_spreadsheet_and_csv_text_cannot_inject_formulas():
    from payroll_reports import safe_text

    assert safe_text('=HYPERLINK("http://malo")') == '\'=HYPERLINK("http://malo")'
    assert safe_text("+57300") == "'+57300" and safe_text("@cmd") == "'@cmd" and safe_text("-1") == "'-1"
    assert safe_text("=SUM(D6:D9)") == "=SUM(D6:D9)" and safe_text("Ana") == "Ana" and safe_text(12) == 12
    evil = run(10)
    evil["lines"][0]["name"] = "=1+1"
    rows = subject.aggregate_benefit([evil], "prima", 2026, 2)
    sheet = load_workbook(BytesIO(subject.build_benefit_workbook("Org", "prima", 2026, 2, rows))).active
    assert sheet["A6"].value == "'=1+1" and sheet["A6"].data_type == "s"
